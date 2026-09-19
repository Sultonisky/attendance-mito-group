"""
Geocode outsource store pin-points from the Excel workbook.

Priority:
1. Google Maps URL coordinates in "pin poin lokasi"
2. Open Location Code (Plus Code) in "Alamat" / pin column
3. Nominatim street-address search (optional, slower)

Writes/merges into data/stores.json for outsource:locations:import.
"""

from __future__ import annotations

import json
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from pathlib import Path

from openlocationcode import openlocationcode as olc
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
EXCEL_PATH = ROOT.parent / "Data titik lokasi kerja karyawan OS.xlsx"
OUTPUT_PATH = ROOT / "stores.json"

CITY_CENTERS = {
    "BALIKOPAN": (-1.2694616, 116.8543264),
    "BANDUNG": (-6.914744, 107.609810),
    "BANGKA": (-2.170, 106.120),
    "BANJAR BARU": (-3.442, 114.842),
    "BANJARMASIN": (-3.318, 114.594),
    "BANYUMAS": (-7.460, 109.287),
    "BANYUWANGI": (-8.219, 114.369),
    "BATAM": (1.130, 104.051),
    "BATU": (-7.866, 112.534),
    "BOGOR": (-6.597, 106.806),
    "MALANG": (-7.983, 112.630),
    "BONTANG": (0.133, 117.481),
    "CIREBON": (-6.706, 108.557),
    "DEPOK": (-6.402, 106.794),
    "GARUT": (-7.227, 107.908),
    "JAKARTA": (-6.2088, 106.8456),
    "JEMBER": (-8.184, 113.668),
    "KEDIRI": (-7.816, 112.016),
    "KETAPANG": (-1.849, 109.971),
    "KLATEN": (-7.705, 110.603),
    "KOTAWARINGIN BARAT": (-2.677, 111.616),
    "KUDUS": (-6.804, 110.840),
    "LAMPUNG-METRO": (-5.117, 105.307),
    "LUBUKLINGGAU": (-3.293, 102.851),
    "MAKASSAR": (-5.147, 119.432),
    "MATARAM": (-8.583, 116.118),
    "MEDAN": (3.595, 98.672),
    "MUARA ENIM": (-3.650, 103.782),
    "PADANG": (-0.949, 100.353),
    "PALEMBANG": (-2.990, 104.757),
    "PALU": (-0.891, 119.870),
    "PASER": (-1.741, 116.427),
    "PATI": (-6.756, 111.036),
    "PEKALONGAN": (-6.888, 109.675),
    "PEKANBARU": (0.507, 101.447),
    "PONOROGO": (-7.867, 111.468),
    "PONTIANAK": (-0.026, 109.342),
    "SAMARINDA": (-0.502, 117.153),
    "SAMPIT": (-2.539, 112.949),
    "SEMARANG": (-6.966, 110.414),
    "SIDOARJO": (-7.4478, 112.7183),
    "SUKOHARJO": (-7.683, 110.84),
    "SURABAYA": (-7.2575, 112.7521),
    "TABALONG": (-2.18, 115.43),
    "TANAH BUMBU": (-3.45, 115.7),
    "TANGERANG": (-6.1783, 106.6319),
    "TENGGARONG": (-0.42, 116.99),
    "TULUNGAGUNG": (-8.0667, 111.9),
    "YOGYAKARTA": (-7.7956, 110.3695),
}

PLUS_CODE_RE = re.compile(
    r"\b([23456789CFGHJMPQRVWX]{4,8}\+[23456789CFGHJMPQRVWX]{2,3})\b",
    re.IGNORECASE,
)
MAPS_AT_RE = re.compile(r"/@(-?\d+\.\d+),(-?\d+\.\d+)")
MAPS_Q_RE = re.compile(r"[?&]q=(-?\d+\.\d+),(-?\d+\.\d+)", re.IGNORECASE)
MAPS_3D_RE = re.compile(r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)")

LOCALITY_ALIASES = {
    "SINGKAWANG": (0.908, 108.987),
    "DENPASAR": (-8.6705, 115.2126),
    "SESETAN": (-8.690, 115.220),
    "TEMANGGUNG": (-7.3167, 110.175),
    "KERTOSARI": (-7.3167, 110.175),
    "MARTAPURA": (-3.411, 114.847),
    "LOKTABAT": (-3.442, 114.842),
    "SEPINGGAN": (-1.268, 116.900),
    "GENTENG KULON": (-8.366, 114.146),
    "ROGOJAMPI": (-8.363, 114.152),
    "PANGKALANBUN": (-2.683, 111.617),
    "PANGKALAN BUN": (-2.683, 111.617),
    "KAMPAR": (0.472, 101.259),
    "TARAI BANGUN": (0.45, 101.40),
    "ANCOL": (-6.125, 106.833),
    "PLUIT": (-6.120, 106.790),
    "KLP. GADING": (-6.157, 106.907),
    "KELAPA GADING": (-6.157, 106.907),
    "BRAGA": (-6.9175, 107.6091),
}

# Verified street-level pins when automated geocoders miss the Excel address.
MANUAL_PINS = {
    ("JEMBER", "TOKO RAMAI JAYA JEMBER"): {
        "lat": -8.1740611,
        "lon": 113.7007280,
        "address": "Jl. R.A. Kartini No.60, Tembaan, Kepatihan, Kec. Kaliwates, Kabupaten Jember, Jawa Timur 68131",
        "source": "excel_manual",
        "confidence": "review",
        "query": "Jalan RA Kartini, Kepatihan, Jember",
    },
}


def normalize_key(city: str, store: str) -> tuple[str, str]:
    return (city.strip().upper(), store.strip().upper())


def city_center(city: str) -> tuple[float, float] | None:
    key = (city or "").strip().upper()
    if not key:
        return None
    if key in CITY_CENTERS:
        return CITY_CENTERS[key]
    compact = re.sub(r"[^A-Z0-9]+", "", key)
    for known, coords in CITY_CENTERS.items():
        known_compact = re.sub(r"[^A-Z0-9]+", "", known)
        if compact == known_compact or compact in known_compact or known_compact in compact:
            return coords
    return None


def reference_from_text(text: str, city: str) -> tuple[float, float]:
    upper = (text or "").upper()
    for alias, coords in LOCALITY_ALIASES.items():
        if alias in upper:
            return coords
    center = city_center(city)
    if center:
        return center
    # last resort: Indonesia centroid — only used when nothing else is known
    return (-2.5489, 118.0149)


def haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    from math import asin, cos, radians, sin, sqrt

    dlat = radians(lat2 - lat1)
    dlon = radians(lon2 - lon1)
    a = sin(dlat / 2) ** 2 + cos(radians(lat1)) * cos(radians(lat2)) * sin(dlon / 2) ** 2
    return 2 * 6371 * asin(sqrt(a))


def extract_maps_coords(text: str | None) -> tuple[float, float] | None:
    if not text:
        return None
    for pattern in (MAPS_AT_RE, MAPS_Q_RE, MAPS_3D_RE):
        match = pattern.search(text)
        if match:
            return float(match.group(1)), float(match.group(2))
    return None


def extract_plus_codes(text: str | None) -> list[str]:
    if not text:
        return []
    seen: set[str] = set()
    out: list[str] = []
    for code in PLUS_CODE_RE.findall(text):
        upper = code.upper()
        if upper not in seen:
            seen.add(upper)
            out.append(upper)
    return out


def locality_snippet(text: str, code: str) -> str:
    """Text after the plus code, useful as Nominatim / recovery context."""
    if not text:
        return ""
    match = re.search(re.escape(code), text, flags=re.IGNORECASE)
    if not match:
        return text[:120]
    tail = text[match.end() :]
    tail = re.sub(r"https?://\S+", " ", tail)
    tail = re.sub(r"\s+", " ", tail).strip(" -;,()")
    return tail[:160]


def decode_plus_code(code: str, city: str, context: str = "") -> tuple[float, float] | None:
    try:
        if olc.isFull(code):
            decoded = olc.decode(code)
            return decoded.latitudeCenter, decoded.longitudeCenter

        candidates: list[tuple[float, float]] = []
        refs = [
            reference_from_text(context, city),
            city_center(city) or (-2.5489, 118.0149),
        ]
        # de-dupe refs
        unique_refs: list[tuple[float, float]] = []
        for ref in refs:
            if ref not in unique_refs:
                unique_refs.append(ref)

        for ref_lat, ref_lon in unique_refs:
            full = olc.recoverNearest(code, ref_lat, ref_lon)
            decoded = olc.decode(full)
            candidates.append((decoded.latitudeCenter, decoded.longitudeCenter))

        city_ref = city_center(city)
        if city_ref:
            # Prefer candidate closest to the branch city, unless locality alias was used
            # and yields a far-but-intentional result (e.g. Denpasar under Surabaya cabang).
            locality_ref = reference_from_text(context, city)
            if locality_ref != city_ref:
                # locality-driven recovery is intentional for mis-tagged branches
                return candidates[0]
            best = min(
                candidates,
                key=lambda c: haversine_km(c[0], c[1], city_ref[0], city_ref[1]),
            )
            if haversine_km(best[0], best[1], city_ref[0], city_ref[1]) <= 250:
                return best
            return None

        return candidates[0] if candidates else None
    except Exception:
        return None


def geocode_pluscode_nominatim(code: str, locality: str, city: str) -> dict | None:
    queries = []
    if locality:
        queries.append(f"{code} {locality}")
    queries.append(f"{code} {city} Indonesia")
    for query in queries:
        results = query_nominatim(query)
        time.sleep(1.1)
        if not results:
            continue
        best = results[0]
        return {
            "lat": float(best["lat"]),
            "lon": float(best["lon"]),
            "address": best.get("display_name") or query,
            "source": "excel_pluscode_nominatim",
            "is_fallback": False,
            "confidence": "high",
            "query": query,
        }
    return None


def clean_street_query(address: str, city: str) -> str:
    text = address
    text = PLUS_CODE_RE.sub(" ", text)
    text = re.sub(r"https?://\S+", " ", text)
    text = re.sub(r"^\s*\d+\.\s*", "", text)
    text = re.sub(r"\s+", " ", text).strip(" -;,")
    if city and city.upper() not in text.upper():
        text = f"{text}, {city}, Indonesia"
    elif "indonesia" not in text.lower():
        text = f"{text}, Indonesia"
    return text


def query_nominatim(query: str) -> list[dict]:
    url = "https://nominatim.openstreetmap.org/search"
    params = {
        "q": query,
        "format": "json",
        "limit": 5,
        "addressdetails": 1,
        "countrycodes": "id",
    }
    headers = {
        "User-Agent": "MITO-Attendance-Geocoder/1.0 (outsource store pinpoints)"
    }
    req = urllib.request.Request(
        f"{url}?{urllib.parse.urlencode(params)}",
        headers=headers,
        method="GET",
    )
    try:
        with urllib.request.urlopen(req, timeout=10) as response:
            data = json.loads(response.read().decode("utf-8"))
        return data if isinstance(data, list) else []
    except (urllib.error.HTTPError, urllib.error.URLError, TimeoutError, json.JSONDecodeError):
        return []


def split_address_candidates(address: str) -> list[str]:
    """Split Excel cells that list multiple numbered outlet addresses."""
    text = (address or "").strip()
    if not text:
        return []
    parts = re.split(r"(?:^|\n)\s*\d+\.\s*", text)
    candidates = [re.sub(r"\s+", " ", p).strip(" -;,") for p in parts if p and p.strip()]
    if not candidates:
        candidates = [re.sub(r"\s+", " ", text).strip()]
    return candidates


def geocode_nominatim(city: str, store: str, address: str) -> dict | None:
    queries: list[str] = []
    for candidate in split_address_candidates(address):
        street = clean_street_query(candidate, city)
        if street:
            queries.append(street)
        simplified = re.sub(r"\bJl\.?\b", "Jalan", candidate, flags=re.IGNORECASE)
        simplified = PLUS_CODE_RE.sub(" ", simplified)
        simplified = re.sub(r"\([^)]*\)", " ", simplified)
        simplified = re.sub(r"\s+", " ", simplified).strip(" -;,")
        if simplified:
            queries.append(simplified)
            queries.append(f"{simplified}, {city}, Indonesia")
            short = re.split(r",\s*", simplified)
            if short:
                queries.append(f"{short[0]}, {city}, Indonesia")
            if len(short) >= 2:
                queries.append(f"{short[0]}, {short[1]}, {city}")
    queries.append(f"{store} {city} Indonesia")

    seen: set[str] = set()
    for query in queries:
        query = re.sub(r"\s+", " ", query).strip()
        if not query or query in seen:
            continue
        seen.add(query)
        results = query_nominatim(query)
        time.sleep(1.1)
        if not results:
            continue
        best = results[0]
        lat = float(best["lat"])
        lon = float(best["lon"])
        city_ref = city_center(city)
        if city_ref and haversine_km(lat, lon, city_ref[0], city_ref[1]) > 80:
            closer = None
            for item in results:
                cand_lat = float(item["lat"])
                cand_lon = float(item["lon"])
                if haversine_km(cand_lat, cand_lon, city_ref[0], city_ref[1]) <= 80:
                    closer = item
                    break
            if closer is None:
                continue
            best = closer
            lat = float(best["lat"])
            lon = float(best["lon"])
        return {
            "lat": lat,
            "lon": lon,
            "address": best.get("display_name") or address,
            "source": "excel_nominatim",
            "is_fallback": False,
            "confidence": "review",
            "query": query,
        }
    return None


def load_excel_stores() -> list[dict]:
    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]

    grouped: dict[tuple[str, str], dict] = {}
    for row in rows:
        city = str(row[1] or "").strip()
        store = str(row[2] or "").strip()
        if not city or not store:
            continue
        key = normalize_key(city, store)
        item = grouped.get(key)
        if item is None:
            item = {
                "city": city.upper() if city.upper() in CITY_CENTERS else city,
                "store": store,
                "alamat": None,
                "pin": None,
                "employees": [],
            }
            # keep canonical city casing from CITY_CENTERS keys when possible
            for known in CITY_CENTERS:
                if known == city.upper():
                    item["city"] = known
                    break
            grouped[key] = item

        employee = str(row[3] or "").strip()
        if employee and employee not in item["employees"]:
            item["employees"].append(employee)
        if row[8] and not item["alamat"]:
            item["alamat"] = str(row[8]).strip()
        if row[7] and not item["pin"]:
            item["pin"] = str(row[7]).strip()

    return list(grouped.values())


def geocode_store(item: dict, use_nominatim: bool = True) -> dict | None:
    city = item["city"]
    store = item["store"]
    alamat = item.get("alamat") or ""
    pin = item.get("pin") or ""
    combined = f"{pin}\n{alamat}".strip()

    maps = extract_maps_coords(pin) or extract_maps_coords(alamat)
    if maps:
        lat, lon = maps
        return {
            "store": store,
            "city": city,
            "lat": lat,
            "lon": lon,
            "address": alamat or pin,
            "source": "excel_maps",
            "is_fallback": False,
            "confidence": "high",
            "query": pin or None,
            "employees": item.get("employees", []),
        }

    plus_codes = extract_plus_codes(combined)
    for code in plus_codes:
        locality = locality_snippet(combined, code)
        decoded = decode_plus_code(code, city, context=combined)
        if decoded:
            lat, lon = decoded
            city_ref = city_center(city)
            locality_ref = reference_from_text(combined, city)
            # If recovery drifted far from both city and locality hints, ask Nominatim.
            drifted = False
            if city_ref and haversine_km(lat, lon, city_ref[0], city_ref[1]) > 250:
                if locality_ref == city_ref or haversine_km(lat, lon, locality_ref[0], locality_ref[1]) > 80:
                    drifted = True
            if not drifted:
                return {
                    "store": store,
                    "city": city,
                    "lat": lat,
                    "lon": lon,
                    "address": alamat or f"{code} ({city})",
                    "source": "excel_pluscode",
                    "is_fallback": False,
                    "confidence": "high",
                    "query": code,
                    "employees": item.get("employees", []),
                }

        if use_nominatim:
            nominatim_pc = geocode_pluscode_nominatim(code, locality, city)
            if nominatim_pc:
                nominatim_pc["store"] = store
                nominatim_pc["city"] = city
                nominatim_pc["employees"] = item.get("employees", [])
                if not nominatim_pc.get("address") and alamat:
                    nominatim_pc["address"] = alamat
                return nominatim_pc

    if use_nominatim and alamat:
        result = geocode_nominatim(city, store, alamat)
        if result:
            result["store"] = store
            result["city"] = city
            result["employees"] = item.get("employees", [])
            return result

    manual = MANUAL_PINS.get(normalize_key(city, store))
    if manual:
        return {
            "store": store,
            "city": city,
            **manual,
            "is_fallback": False,
            "employees": item.get("employees", []),
        }

    return None


def should_replace(existing: dict, incoming: dict) -> bool:
    """Prefer Excel pinpoints over city fallbacks / weaker sources."""
    source_rank = {
        "excel_maps": 100,
        "excel_pluscode": 95,
        "excel_pluscode_nominatim": 90,
        "excel_manual": 85,
        "excel_nominatim": 70,
        "nominatim": 50,
        "city_fallback": 10,
    }
    existing_rank = source_rank.get(existing.get("source", ""), 40)
    incoming_rank = source_rank.get(incoming.get("source", ""), 40)

    if existing.get("is_fallback"):
        return True
    if incoming_rank > existing_rank:
        return True
    if incoming_rank == existing_rank and incoming.get("source", "").startswith("excel"):
        return True
    return False


def merge_into_stores(existing: list[dict], geocoded: list[dict]) -> tuple[list[dict], dict]:
    index = {
        normalize_key(str(item.get("city", "")), str(item.get("store", ""))): i
        for i, item in enumerate(existing)
    }
    stats = {
        "updated": 0,
        "added": 0,
        "kept": 0,
        "geocoded": len(geocoded),
        "by_source": defaultdict(int),
    }

    for record in geocoded:
        key = normalize_key(record["city"], record["store"])
        stats["by_source"][record["source"]] += 1
        # Keep employees only in a side field for traceability; importer ignores unknown keys.
        payload = {
            "store": record["store"],
            "city": record["city"],
            "lat": record["lat"],
            "lon": record["lon"],
            "address": record["address"],
            "source": record["source"],
            "is_fallback": False,
            "confidence": record.get("confidence", "review"),
            "query": record.get("query"),
        }
        if record.get("employees"):
            payload["employees"] = record["employees"]

        if key in index:
            current = existing[index[key]]
            if should_replace(current, payload):
                existing[index[key]] = payload
                stats["updated"] += 1
            else:
                # still attach employees if missing
                if "employees" not in current and record.get("employees"):
                    current["employees"] = record["employees"]
                stats["kept"] += 1
        else:
            existing.append(payload)
            index[key] = len(existing) - 1
            stats["added"] += 1

    return existing, stats


def main() -> None:
    print(f"Reading Excel: {EXCEL_PATH}")
    excel_stores = load_excel_stores()
    print(f"Unique city+store from Excel: {len(excel_stores)}")

    with_hint = [
        item
        for item in excel_stores
        if item.get("alamat") or item.get("pin")
    ]
    print(f"Stores with alamat/pin hint: {len(with_hint)}")

    geocoded: list[dict] = []
    failed: list[str] = []

    for idx, item in enumerate(with_hint, start=1):
        label = f"{item['city']} / {item['store']}"
        result = geocode_store(item, use_nominatim=True)
        if result:
            geocoded.append(result)
            print(
                f"[{idx}/{len(with_hint)}] OK {label} "
                f"=> {result['lat']:.6f},{result['lon']:.6f} ({result['source']})"
            )
        else:
            failed.append(label)
            print(f"[{idx}/{len(with_hint)}] FAIL {label}")

    existing: list[dict] = []
    if OUTPUT_PATH.exists():
        try:
            loaded = json.loads(OUTPUT_PATH.read_text(encoding="utf-8"))
            if isinstance(loaded, list):
                existing = loaded
        except json.JSONDecodeError:
            existing = []

    merged, stats = merge_into_stores(existing, geocoded)
    OUTPUT_PATH.write_text(
        json.dumps(merged, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    real = sum(1 for item in merged if not item.get("is_fallback"))
    fallback = sum(1 for item in merged if item.get("is_fallback"))
    print("\n=== Summary ===")
    print(f"Geocoded from Excel: {stats['geocoded']}")
    print(f"Updated existing rows: {stats['updated']}")
    print(f"Added new rows: {stats['added']}")
    print(f"Kept stronger existing: {stats['kept']}")
    print(f"By source: {dict(stats['by_source'])}")
    print(f"stores.json total={len(merged)} real={real} fallback={fallback}")
    if failed:
        print("Failed:")
        for label in failed:
            print(f"  - {label}")


if __name__ == "__main__":
    main()

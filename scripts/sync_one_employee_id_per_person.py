#!/usr/bin/env python3
"""Sync Excel EMPLOYEE ID: 1 person = 1 sequential code (DM20260001..N).

1) Collapse duplicate IDs per NAMA KARYAWAN (keep first appearance order)
2) Renumber sequentially DM20260001, DM20260002, ...
3) Regenerate data/stores.json
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "Data titik lokasi kerja karyawan OS (FIXED).xlsx"
REPORT = ROOT / "data" / "employee_id_collapse_report.json"


def norm_name(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip()).casefold()


def main() -> int:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col = {h: i + 1 for i, h in enumerate(headers)}
    for need in ("NAMA KARYAWAN", "EMPLOYEE ID"):
        if need not in col:
            raise SystemExit(f"Missing column {need}")

    # First pass: discover people in sheet order
    people: OrderedDict[str, dict] = OrderedDict()
    row_people: list[tuple[int, str, str]] = []  # row, name_key, old_id

    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, col["NAMA KARYAWAN"]).value
        eid = ws.cell(r, col["EMPLOYEE ID"]).value
        if name is None and eid is None:
            continue
        key = norm_name(name)
        eid_s = str(eid).strip() if eid is not None else ""
        if not key:
            continue
        row_people.append((r, key, eid_s))
        if key not in people:
            people[key] = {
                "employee": str(name).strip(),
                "old_first_id": eid_s,
                "first_row": r,
            }

    # Sequential codes
    seq_map: dict[str, str] = {}
    for i, key in enumerate(people.keys(), start=1):
        seq_map[key] = f"DM2026{i:04d}"

    changes = []
    for r, key, old_id in row_people:
        new_id = seq_map[key]
        if old_id != new_id:
            changes.append(
                {
                    "source_row": r,
                    "employee": people[key]["employee"],
                    "old_employee_id": old_id,
                    "new_employee_id": new_id,
                }
            )
        ws.cell(r, col["EMPLOYEE ID"]).value = new_id

    wb.save(XLSX)

    REPORT.write_text(
        json.dumps(
            {
                "people": len(people),
                "sequence_from": "DM20260001",
                "sequence_to": f"DM2026{len(people):04d}",
                "rows_rewritten": len(changes),
                "changes": changes,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )

    print(f"people={len(people)} sequence=DM20260001..DM2026{len(people):04d}")
    print(f"rows_rewritten={len(changes)}")
    print(f"wrote {XLSX}")
    print(f"report {REPORT}")

    rc = subprocess.call([sys.executable, str(ROOT / "scripts" / "normalize_outsource_os_locations.py")])
    if rc != 0:
        return rc

    # Verify contiguous sequence in JSON
    stores = json.loads((ROOT / "data" / "stores.json").read_text(encoding="utf-8"))
    ids = []
    seen = set()
    for row in sorted(stores, key=lambda x: x["source_row"]):
        eid = row["employee_id"]
        if eid not in seen:
            seen.add(eid)
            ids.append(eid)
    expected = [f"DM2026{i:04d}" for i in range(1, len(ids) + 1)]
    if ids != expected:
        print("FAIL: employee_id sequence not contiguous")
        print(" got:", ids[:15], "...")
        print(" want:", expected[:15], "...")
        return 1
    print(f"stores.json OK: {len(ids)} sequential employee_ids")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

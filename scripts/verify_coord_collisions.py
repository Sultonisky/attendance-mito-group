#!/usr/bin/env python3
"""Validate stores.json pins: no city-fallback flags; flag only bad geocode collisions."""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "Data titik lokasi kerja karyawan OS (FIXED).xlsx"
JSON_PATH = ROOT / "data" / "stores.json"

PLUS_RE = re.compile(
    r"\b([2-9CFGHJMPQRVWX]{2,8}\+[2-9CFGHJMPQRVWX]{2,3})\b",
    re.I,
)


def norm_addr(a: str) -> str:
    s = re.sub(r"\([^)]*\)", " ", a or "")
    s = re.sub(r"^[^:]{0,40}:\s*", "", s)
    if "=" in s:
        s = s.split("=", 1)[-1]
    s = re.sub(r"\s+", " ", s).strip().lower()
    return s


def street_key(a: str) -> str:
    s = norm_addr(a)
    # first segment before comma, strip no.
    first = s.split(",")[0]
    first = re.sub(r"\bno\.?\s*\d+\w*", "", first)
    return re.sub(r"\s+", " ", first).strip()


def plus_codes(a: str) -> set[str]:
    return {m.group(1).upper() for m in PLUS_RE.finditer(a or "")}


def main() -> int:
    stores = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    method_col = headers.index("METHOD") + 1

    def method_of(sr: int) -> str:
        return str(ws.cell(sr, method_col).value or "")

    errors: list[str] = []
    if any(s.get("is_fallback") for s in stores):
        errors.append("is_fallback present")
    if any(s.get("source") != "excel_normalized" for s in stores):
        errors.append("non-excel_normalized source")
    if any(s.get("confidence") != "high" for s in stores):
        errors.append("confidence not high")
    if any(s.get("incomplete_reason") for s in stores):
        errors.append("incomplete rows")
    if any(s.get("lat") is None or s.get("lon") is None for s in stores):
        errors.append("null coords")
    if any(not (s.get("address") or "").strip() for s in stores):
        errors.append("empty address")
    if any(not (-11.5 <= s["lat"] <= 6.5 and 94.5 <= s["lon"] <= 141.5) for s in stores):
        errors.append("outside Indonesia bbox")

    groups: dict[tuple[float, float], list[dict]] = defaultdict(list)
    for s in stores:
        groups[(round(float(s["lat"]), 5), round(float(s["lon"]), 5))].append(s)

    bad_collisions = []
    ok_shared = []
    for key, rows in groups.items():
        if len(rows) < 2:
            continue
        addrs = {norm_addr(r["address"]) for r in rows}
        pluses = set()
        for r in rows:
            pluses |= plus_codes(r["address"])
        streets = {street_key(r["address"]) for r in rows if street_key(r["address"])}

        # Legitimate share: identical address, shared plus-code, or same street stem
        if len(addrs) == 1 or (len(pluses) == 1 and len(pluses) > 0) or (len(streets) == 1 and len(streets) > 0):
            ok_shared.append({"coord": key, "n": len(rows), "reason": "same_pin_or_plus_or_street"})
            continue

        bad_collisions.append(
            {
                "coord": key,
                "n": len(rows),
                "employee_ids": [r["employee_id"] for r in rows],
                "methods": [method_of(r["source_row"]) for r in rows],
            }
        )

    print("=== FLAG CHECK ===")
    print(f"rows={len(stores)}")
    print(f"is_fallback_true={sum(1 for s in stores if s.get('is_fallback'))}")
    print(f"confidence={Counter(s.get('confidence') for s in stores)}")
    print(f"incomplete={sum(1 for s in stores if s.get('incomplete_reason'))}")
    print(f"excel_METHOD={Counter(method_of(s['source_row']) for s in stores)}")
    print(f"unique_coords={len(groups)}")
    print(f"ok_shared_groups={len(ok_shared)}")
    print(f"bad_collision_groups={len(bad_collisions)}")
    for c in bad_collisions:
        print(" BAD", c)

    print("\n=== ERRORS ===")
    print(f"errors={len(errors)}")
    for e in errors:
        print(" -", e)

    if errors or bad_collisions:
        print("\nRESULT: FAIL")
        return 1

    print("\nRESULT: PASS — all pins valid; shared coords are same plus-code/street/address only")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
